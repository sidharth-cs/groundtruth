"""Reading documents off disk and turning them into addressable chunks.

Two things matter here and nothing else does.

First, offsets must be exact. Every citation the system emits later is a pair
of character offsets into the normalised text produced here. If normalisation
and chunking disagree by a single character, every citation in the register
becomes unverifiable. `Chunk` enforces `len(text) == end - start`, so a drift
fails loudly at construction rather than quietly producing a register full of
plausible-looking nonsense.

Second, parsing is per-format but normalisation is shared. Whatever comes in —
plain text, Markdown, HTML, DOCX, PDF — becomes one normalised string, and
from that point the rest of the system does not care what the file was.
"""

from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path
from uuid import UUID

from app.domain.models import Chunk, Document, DocumentStatus

# Formats this system accepts. Stated here, and in the README, because the
# brief asks for the declared set — a second run means different documents
# inside this set, not a different set.
SUPPORTED_SUFFIXES = {".txt", ".md", ".html", ".htm", ".docx", ".pdf"}

_MEDIA_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
}


class _TextExtractor(HTMLParser):
    """Strip tags, keep text, insert breaks at block boundaries.

    Deliberately not a full HTML renderer. It only needs to produce a stable
    plain-text projection that chunking can index into.
    """

    _BLOCK = {
        "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
        "section", "article", "table", "blockquote",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in {"script", "style"}:
            self._skip += 1
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self._skip = max(0, self._skip - 1)
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self._parts.append(data)

    def text(self) -> str:
        return "".join(self._parts)


def _read_html(path: Path) -> str:
    parser = _TextExtractor()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    parser.close()
    return parser.text()


def _read_docx(path: Path) -> str:
    from docx import Document as Docx

    doc = Docx(str(path))
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables carry real content in this domain (invoice lines, ownership
    # tables), so they are flattened into the text rather than dropped.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    # Joined with a blank line so each Word paragraph becomes its own chunk.
    # Joining with a single newline collapses the whole file into one chunk,
    # and a citation that points at an entire contract is not evidence a
    # reviewer can check.
    return "\n\n".join(blocks)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def normalise(raw: str) -> str:
    """Collapse whitespace into a stable, indexable form.

    Runs of blank lines become exactly one blank line, trailing spaces go, and
    line endings are unified. This happens once, before chunking, so offsets
    are taken against the same string the citations will be checked against.
    """
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_document(path: Path, pile_id: str) -> Document:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"{path.name}: unsupported format {suffix!r}. "
            f"Accepted: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    fail_reason: str | None = None
    try:
        if suffix in {".html", ".htm"}:
            raw = _read_html(path)
        elif suffix == ".docx":
            raw = _read_docx(path)
        elif suffix == ".pdf":
            raw = _read_pdf(path)
        else:
            raw = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        # A parser failure on one file must not take down the run — the file
        # becomes FAILED, same as the empty-text case below, and the pile
        # keeps going. Isolating this per-document is the point; the run
        # continuing while naming the file it could not read is what turns a
        # crash into an honest, recoverable result.
        raw = ""
        fail_reason = f"{path.name}: could not parse as {suffix}: {exc}"

    text = normalise(raw)
    status = DocumentStatus.PARSED if text else DocumentStatus.FAILED
    return Document(
        pile_id=pile_id,
        filename=path.name,
        media_type=_MEDIA_TYPES[suffix],
        text=text,
        status=status,
        quarantine_reason=fail_reason if status is DocumentStatus.FAILED else None,
    )


def chunk_document(document: Document) -> list[Chunk]:
    """Split on blank lines, preserving exact offsets.

    Paragraph-level granularity is the right unit for this domain: a claim
    ("shareholding 62%") lives inside one paragraph, and a citation that
    points at a whole page is not evidence a reviewer can check quickly.
    """
    chunks: list[Chunk] = []
    ordinal = 0
    for match in re.finditer(r"[^\n]+(?:\n[^\n]+)*", document.text):
        body = match.group(0)
        if not body.strip():
            continue
        chunks.append(
            Chunk(
                document_id=document.id,
                ordinal=ordinal,
                text=body,
                char_start=match.start(),
                char_end=match.end(),
            )
        )
        ordinal += 1
    return chunks


def load_pile(directory: Path, pile_id: str) -> list[tuple[Document, list[Chunk]]]:
    """Read every supported document in a directory, in stable filename order."""
    if not directory.is_dir():
        raise FileNotFoundError(f"no such pile directory: {directory}")

    loaded: list[tuple[Document, list[Chunk]]] = []
    for path in sorted(directory.iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        document = read_document(path, pile_id)
        loaded.append((document, chunk_document(document)))
    return loaded


def slice_for(document: Document, char_start: int, char_end: int) -> str:
    """Re-slice the source. Used to verify a citation actually holds."""
    return document.text[char_start:char_end]


def find_chunk(chunks: list[Chunk], char_start: int) -> UUID | None:
    """Which chunk contains this offset."""
    for chunk in chunks:
        if chunk.char_start <= char_start < chunk.char_end:
            return chunk.id
    return None
